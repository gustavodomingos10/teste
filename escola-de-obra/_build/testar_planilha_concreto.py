#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Testa a planilha do concreto: injeta 4 caminhões + 2 CPs e avalia as fórmulas (lib 'formulas')."""
import os, re, sys, shutil, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
import formulas

BASE=os.path.abspath(os.path.join(os.path.dirname(__file__),".."))
SRC=os.path.join(BASE,"05-curso-concreto","kit","planilha","planilha-concreto.xlsx")
TMP="/tmp/planilha-concreto-teste.xlsx"
shutil.copy(SRC,TMP)

# Parametros: slump 100±20mm, limite 150min, fck 30
# horários como fração do dia: 10:00=10/24, fim conforme caso
caminhoes=[
 # NF, hora_agua, fim_descarga, slump, nota, aspecto, agua, esperado(veredito, tempo_min)
 ("NF-1", 10/24, 11.5/24, 105, "Sim","Sim","Nao", ("LIBERADO", 90)),
 ("NF-2", 10/24, 13/24,   100, "Sim","Sim","Nao", ("RECUSAR - tempo", 180)),
 ("NF-3", 10/24, 11/24,    70, "Sim","Sim","Nao", ("CORRIGIR VIA CENTRAL", 60)),
 ("NF-4", 10/24, 11/24,   150, "Sim","Sim","Nao", ("RECUSAR - fluido", 60)),
]
cps=[("CP-1","NF-1",32.5,"OK"),("CP-2","NF-1",26.0,"ALERTA")]

wb=load_workbook(TMP)
ws=wb["Recebimento"]
for i,(nf,ha,fd,sl,nota,asp,agua,_) in enumerate(caminhoes):
    r=3+i
    ws[f"B{r}"]=nf; ws[f"C{r}"]=ha; ws[f"D{r}"]=fd; ws[f"F{r}"]=sl
    ws[f"G{r}"]=nota; ws[f"H{r}"]=asp; ws[f"I{r}"]=agua
wc=wb["CPs"]
for i,(cid,nf,res,_) in enumerate(cps):
    r=3+i
    wc[f"A{r}"]=cid; wc[f"B{r}"]=nf; wc[f"F{r}"]=res
wb.save(TMP)

xl=formulas.ExcelModel().loads(TMP).finish(); sol=xl.calculate()
def get(sheet,cell):
    pat=re.compile(r"\]"+sheet.upper()+r"'?\!"+cell.upper()+r"$")
    for k,v in sol.items():
        if pat.search(k):
            try: val=v.value[0,0]
            except Exception:
                try: val=v.value
                except Exception: val=v
            if hasattr(val,"ravel"): val=val.ravel()[0]
            return val
    return None

ok=True
print("="*64); print(" TESTE DA PLANILHA DE CONCRETO"); print("="*64)
for i,(nf,ha,fd,sl,nota,asp,agua,esp) in enumerate(caminhoes):
    r=3+i
    ver=str(get("Recebimento",f"J{r}")).strip().strip('"')
    tempo=get("Recebimento",f"E{r}")
    try: tempo=int(round(float(tempo)))
    except Exception: pass
    e_ver,e_t=esp
    v_ok=ver==e_ver; t_ok=tempo==e_t
    passed=v_ok and t_ok; ok=ok and passed
    print(f"\n{'✔' if passed else '✗'} {nf}: slump={sl}mm tempo={tempo}min (esp {e_t})")
    print(f"   Veredito: {ver}  (esperado {e_ver}) {'✓' if v_ok else '✗'}")
for i,(cid,nf,res,e_pref) in enumerate(cps):
    r=3+i
    sit=str(get("CPs",f"G{r}")).strip().strip('"')
    s_ok=sit.startswith(e_pref); ok=ok and s_ok
    print(f"\n{'✔' if s_ok else '✗'} {cid}: {res} MPa → {sit}")

def num(x):
    try: return int(round(float(x)))
    except Exception: return x
cam=num(get("Dashboard","A4")); lib=num(get("Dashboard","B4")); cor=num(get("Dashboard","C4")); rec=num(get("Dashboard","D4"))
cpn=num(get("Dashboard","A7")); cpo=num(get("Dashboard","B7")); cpa=num(get("Dashboard","C7"))
print("\n"+"-"*64); print(" DASHBOARD")
dash_ok=(cam==4 and lib==1 and cor==1 and rec==2 and cpn==2 and cpo==1 and cpa==1)
print(f"  Caminhões={cam}(4) Liberados={lib}(1) Corrigir={cor}(1) Recusados={rec}(2)")
print(f"  CPs={cpn}(2) OK={cpo}(1) Alerta={cpa}(1)  {'✔' if dash_ok else '✗'}")
ok=ok and dash_ok
print("\n"+"="*64); print(" RESULTADO FINAL:", "TODOS OS TESTES PASSARAM ✔" if ok else "HÁ FALHAS ✗"); print("="*64)
sys.exit(0 if ok else 1)
