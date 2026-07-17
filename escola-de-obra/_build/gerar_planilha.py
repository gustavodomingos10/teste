#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera a Planilha Inteligente de Monitoramento de Fissuras (Kit Diamante).
Abas: Parametros, Lancamentos, Dashboard.
Fórmulas em sintaxe internacional (separador vírgula) — o Excel PT-BR converte a exibição.
Marca: verde-petróleo / creme / dourado (brand-tokens.md).
Saída: 02-curso-piloto/kit/planilha/planilha-fissuras.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

# ---- Cores da marca ----
VERDE = "0E3A34"; VERDE_MED = "1C5A4E"; CREME = "F5EFE0"; CREME2 = "FBF8F0"
DOURADO = "C9A24B"; CARVAO = "20211E"; CINZA = "6B6B63"
V_VERDE = "2E7D32"; V_AMAR = "F9A825"; V_VERM = "C62828"

f_titulo = Font(name="Georgia", size=15, bold=True, color=CREME)
f_hdr = Font(name="Calibri", size=10, bold=True, color=CREME)
f_lbl = Font(name="Calibri", size=10, bold=True, color=VERDE)
f_norm = Font(name="Calibri", size=10, color=CARVAO)
f_small = Font(name="Calibri", size=8.5, italic=True, color=CINZA)
fill_verde = PatternFill("solid", fgColor=VERDE)
fill_vmed = PatternFill("solid", fgColor=VERDE_MED)
fill_creme = PatternFill("solid", fgColor=CREME)
fill_creme2 = PatternFill("solid", fgColor=CREME2)
fill_dour = PatternFill("solid", fgColor=DOURADO)
fill_alta = PatternFill("solid", fgColor="F4CCCC")
fill_media = PatternFill("solid", fgColor="FFF2CC")
fill_baixa = PatternFill("solid", fgColor="D9EAD3")
thin = Side(style="thin", color="D9D2BE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()

# ============================================================
# ABA 1 — PARAMETROS
# ============================================================
wsp = wb.active
wsp.title = "Parametros"
wsp.sheet_view.showGridLines = False
wsp["A1"] = "PARÂMETROS DE DECISÃO — ajuste e assine antes de usar"
wsp["A1"].font = f_titulo; wsp["A1"].fill = fill_verde
wsp.merge_cells("A1:D1"); wsp.row_dimensions[1].height = 26
wsp["A1"].alignment = Alignment(horizontal="left", vertical="center")

params = [
    ("Limiar de atividade (mm)", 0.10, "Crescimento de abertura que caracteriza fissura ATIVA no período. [VALIDAR B9]"),
    ("Limite de abertura w_k (mm)", 0.30, "Abertura de referência (NBR 6118, Tab. 13.4 — CAA II/III). Ajustar à classe de agressividade. [VALIDAR B1]"),
]
wsp["A3"] = "Parâmetro"; wsp["B3"] = "Valor"; wsp["C3"] = "Observação / fonte"
for c in ("A3","B3","C3"):
    wsp[c].font = f_hdr; wsp[c].fill = fill_vmed; wsp[c].alignment = left; wsp[c].border = border
r = 4
for nome, val, obs in params:
    wsp[f"A{r}"] = nome; wsp[f"A{r}"].font = f_lbl
    wsp[f"B{r}"] = val; wsp[f"B{r}"].font = f_norm; wsp[f"B{r}"].number_format = "0.00"
    wsp[f"B{r}"].fill = fill_dour; wsp[f"B{r}"].alignment = center
    wsp[f"C{r}"] = obs; wsp[f"C{r}"].font = f_small; wsp[f"C{r}"].alignment = left
    for col in "ABC": wsp[f"{col}{r}"].border = border
    r += 1

# Listas para validação (documentadas)
listas = {
    "Orientacoes": ["Vertical","Horizontal","Diagonal 45","Mapeada"],
    "Locais": ["Meio do vao","Apoio da viga","Canto de abertura","Base da parede","Topo da parede","Pilar","Laje","Muro"],
    "Origens": ["Flexao","Cisalhamento","Recalque diferencial","Retracao","Expansao alvenaria","Umidade","Concentracao de tensao (canto)"],
    "SimNao": ["Sim","Nao"],
}
wsp["A8"] = "Listas de referência (usadas nos menus suspensos de Lançamentos)"
wsp["A8"].font = f_lbl; wsp.merge_cells("A8:D8")
col = 1
for nome, vals in listas.items():
    letter = get_column_letter(col)
    wsp[f"{letter}9"] = nome; wsp[f"{letter}9"].font = f_hdr; wsp[f"{letter}9"].fill = fill_vmed
    wsp[f"{letter}9"].border = border; wsp[f"{letter}9"].alignment = center
    for i, v in enumerate(vals):
        cell = wsp[f"{letter}{10+i}"]; cell.value = v; cell.font = f_norm; cell.border = border
    col += 1
wsp.column_dimensions["A"].width = 26; wsp.column_dimensions["B"].width = 14
wsp.column_dimensions["C"].width = 40; wsp.column_dimensions["D"].width = 22

# ============================================================
# ABA 2 — LANCAMENTOS
# ============================================================
wsl = wb.create_sheet("Lancamentos")
wsl.sheet_view.showGridLines = False
wsl["A1"] = "MONITORAMENTO DE FISSURAS — uma linha por fissura numerada"
wsl["A1"].font = f_titulo; wsl["A1"].fill = fill_verde
wsl.merge_cells("A1:P1"); wsl.row_dimensions[1].height = 24
wsl["A1"].alignment = Alignment(horizontal="left", vertical="center")

headers = [
    ("A","Nº",7),("B","Ambiente / Elemento",20),("C","Orientação",13),("D","Local no elemento",16),
    ("E","Abertura inicial (mm)",11),("F","Data 1ª medição",13),("G","Abertura atual (mm)",11),
    ("H","Data última medição",13),("I","Δ abertura (mm)",11),("J","Situação",11),
    ("K","Origem (hipótese)",18),("L","Umidade?",9),("M","Risco imediato?",11),
    ("N","Gravidade",11),("O","Encaminhamento",30),("P","Observações",26),
]
for col, txt, w in headers:
    c = wsl[f"{col}2"]; c.value = txt; c.font = f_hdr; c.fill = fill_verde
    c.alignment = center; c.border = border
    wsl.column_dimensions[col].width = w
wsl.row_dimensions[2].height = 30
wsl.freeze_panes = "A3"

FIRST, LAST = 3, 202  # linhas de dados
for row in range(FIRST, LAST+1):
    wsl[f"I{row}"] = f'=IF(G{row}="","",G{row}-E{row})'
    wsl[f"J{row}"] = (f'=IF(G{row}="","",IF(ABS(I{row})>=Parametros!$B$4,"ATIVA","ESTAVEL"))')
    wsl[f"N{row}"] = (
        f'=IF(G{row}="","",'
        f'IF(M{row}="Sim","ALTA",'
        f'IF(AND(OR(K{row}="Cisalhamento",K{row}="Recalque diferencial"),J{row}="ATIVA"),"ALTA",'
        f'IF(OR(J{row}="ATIVA",G{row}>Parametros!$B$5),"MEDIA","BAIXA"))))'
    )
    wsl[f"O{row}"] = (
        f'=IF(N{row}="","",'
        f'IF(N{row}="ALTA","Escorar/interditar e acionar calculista",'
        f'IF(N{row}="MEDIA","Tratar a causa + monitorar","Monitorar e registrar")))'
    )
    for col, _, _ in headers:
        cell = wsl[f"{col}{row}"]
        cell.font = f_norm; cell.border = border
        cell.alignment = center if col in ("A","C","D","E","F","G","H","I","J","L","M","N") else left
        if row % 2 == 0: cell.fill = fill_creme2
    wsl[f"E{row}"].number_format = "0.00"; wsl[f"G{row}"].number_format = "0.00"
    wsl[f"I{row}"].number_format = "0.00"
    wsl[f"F{row}"].number_format = "dd/mm/yyyy"; wsl[f"H{row}"].number_format = "dd/mm/yyyy"

# Validação de dados (menus suspensos)
dv_map = {
    "C": '"Vertical,Horizontal,Diagonal 45,Mapeada"',
    "D": '"Meio do vao,Apoio da viga,Canto de abertura,Base da parede,Topo da parede,Pilar,Laje,Muro"',
    "K": '"Flexao,Cisalhamento,Recalque diferencial,Retracao,Expansao alvenaria,Umidade,Concentracao de tensao (canto)"',
    "L": '"Sim,Nao"', "M": '"Sim,Nao"',
}
for col, formula in dv_map.items():
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Escolha um valor da lista."; dv.errorTitle = "Valor inválido"
    dv.prompt = "Selecione uma opção"; dv.showInputMessage = True
    wsl.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")

# Formatação condicional — semáforo na coluna Gravidade (N)
rng = f"N{FIRST}:N{LAST}"
wsl.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ALTA"'], fill=fill_alta, font=Font(color=V_VERM, bold=True)))
wsl.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"MEDIA"'], fill=fill_media, font=Font(color="9C6500", bold=True)))
wsl.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"BAIXA"'], fill=fill_baixa, font=Font(color=V_VERDE, bold=True)))
# Situação ATIVA em destaque
wsl.conditional_formatting.add(f"J{FIRST}:J{LAST}", CellIsRule(operator="equal", formula=['"ATIVA"'], fill=fill_media, font=Font(color="9C6500", bold=True)))

# ============================================================
# ABA 3 — DASHBOARD
# ============================================================
wsd = wb.create_sheet("Dashboard")
wsd.sheet_view.showGridLines = False
wsd["A1"] = "DASHBOARD — Painel de Monitoramento de Fissuras"
wsd["A1"].font = f_titulo; wsd["A1"].fill = fill_verde
wsd.merge_cells("A1:F1"); wsd.row_dimensions[1].height = 28
wsd["A1"].alignment = Alignment(horizontal="left", vertical="center")

def metric(cell_lbl, cell_val, label, formula, fmt="0"):
    wsd[cell_lbl] = label; wsd[cell_lbl].font = f_lbl; wsd[cell_lbl].alignment = left
    wsd[cell_val] = formula; wsd[cell_val].font = Font(name="Georgia", size=18, bold=True, color=VERDE)
    wsd[cell_val].alignment = center; wsd[cell_val].fill = fill_creme2; wsd[cell_val].number_format = fmt
    wsd[cell_lbl].border = border; wsd[cell_val].border = border

metric("A3","A4","Total de fissuras","=COUNT(Lancamentos!A3:A202)")
metric("B3","B4","Ativas","=COUNTIF(Lancamentos!J3:J202,\"ATIVA\")")
metric("C3","C4","Estáveis","=COUNTIF(Lancamentos!J3:J202,\"ESTAVEL\")")

# Tabela de gravidade (base do gráfico)
wsd["A7"] = "Gravidade"; wsd["B7"] = "Quantidade"
for c in ("A7","B7"): wsd[c].font = f_hdr; wsd[c].fill = fill_vmed; wsd[c].border = border; wsd[c].alignment = center
grav = [("ALTA",'=COUNTIF(Lancamentos!N3:N202,"ALTA")',fill_alta),
        ("MEDIA",'=COUNTIF(Lancamentos!N3:N202,"MEDIA")',fill_media),
        ("BAIXA",'=COUNTIF(Lancamentos!N3:N202,"BAIXA")',fill_baixa)]
for i,(lbl,fml,fl) in enumerate(grav):
    wsd[f"A{8+i}"] = lbl; wsd[f"A{8+i}"].fill = fl; wsd[f"A{8+i}"].font = Font(bold=True, color=CARVAO)
    wsd[f"A{8+i}"].border = border; wsd[f"A{8+i}"].alignment = center
    wsd[f"B{8+i}"] = fml; wsd[f"B{8+i}"].font = f_norm; wsd[f"B{8+i}"].border = border; wsd[f"B{8+i}"].alignment = center

# Alerta condicional
wsd["A12"] = "Alerta"
wsd["A12"].font = f_hdr; wsd["A12"].fill = fill_vmed; wsd["A12"].border = border; wsd["A12"].alignment = center
wsd["B12"] = ('=IF(B8>0,"HA FISSURAS DE GRAVIDADE ALTA - priorize o fluxograma de emergencia e comunique por escrito",'
              'IF(B9>0,"Atencao: fissuras de gravidade media em monitoramento","Sem gravidade alta/media no momento"))')
wsd.merge_cells("B12:F12"); wsd["B12"].font = Font(bold=True, color=V_VERM)
wsd["B12"].alignment = left; wsd["B12"].border = border

# Gráfico de barras
chart = BarChart(); chart.type = "col"; chart.title = "Fissuras por gravidade"
chart.style = 10; chart.y_axis.title = "Qtde"; chart.x_axis.title = "Gravidade"
data = Reference(wsd, min_col=2, min_row=7, max_row=10)
cats = Reference(wsd, min_col=1, min_row=8, max_row=10)
chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
chart.height = 7.5; chart.width = 13; chart.legend = None
wsd.add_chart(chart, "D7")

wsd["A15"] = ("Como usar: preencha a aba Lancamentos (uma linha por fissura). Situacao, Gravidade e "
              "Encaminhamento sao calculados automaticamente a partir dos Parametros. Ajuste os limiares "
              "na aba Parametros e assine.")
wsd["A15"].font = f_small; wsd.merge_cells("A15:F16"); wsd["A15"].alignment = left
wsd["A18"] = "Revisado e assinado por Gustavo Domingos — Engenheiro Civil, CREA-PR 140.964-D"
wsd["A18"].font = Font(name="Calibri", size=9, bold=True, color=DOURADO)
for w,c in [(14,"A"),(12,"B"),(12,"C"),(14,"D"),(12,"E"),(12,"F")]:
    wsd.column_dimensions[c].width = w

out = os.path.join(os.path.dirname(__file__), "..", "02-curso-piloto", "kit", "planilha", "planilha-fissuras.xlsx")
out = os.path.abspath(out)
wb.save(out)
print("PLANILHA GERADA:", out)
