#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Termo de Liberação de Concretagem (.docx) — Kit do Curso 4."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERDE=RGBColor(0x0E,0x3A,0x34); DOUR=RGBColor(0xC9,0xA2,0x4B); CINZA=RGBColor(0x6B,0x6B,0x63); REALCE=RGBColor(0x9C,0x65,0x00)
doc=Document(); st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexc); tcPr.append(sh)
def h(txt,size=13,sb=10):
    p=doc.add_paragraph(); r=p.add_run(txt); r.bold=True; r.font.size=Pt(size); r.font.color.rgb=VERDE; r.font.name="Georgia"
    p.paragraph_format.space_before=Pt(sb); p.paragraph_format.space_after=Pt(3); return p
def campo(label,hint):
    p=doc.add_paragraph(); r=p.add_run(label+" "); r.bold=True; r.font.color.rgb=VERDE
    r2=p.add_run("["+hint+"]"); r2.font.color.rgb=REALCE; r2.italic=True; return p
def tabela(cols, rows, fillhdr="0E3A34"):
    t=doc.add_table(rows=1,cols=len(cols)); t.style="Table Grid"
    for i,c in enumerate(cols):
        cell=t.rows[0].cells[i]; shade(cell,fillhdr)
        rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xF5,0xEF,0xE0); rr.font.size=Pt(9)
    for _ in range(rows):
        cells=t.add_row().cells
        for i in range(len(cols)):
            rr=cells[i].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(9)
    return t

tt=doc.add_paragraph(); tt.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tt.add_run("TERMO DE LIBERAÇÃO DE CONCRETAGEM"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=VERDE; r.font.name="Georgia"
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
rs=sub.add_run("Escola de Obra · GD Engenharia e Perícia · Padrão Diamante"); rs.font.color.rgb=DOUR; rs.bold=True; rs.font.size=Pt(9)
nota=doc.add_paragraph()
rn=nota.add_run("COMO USAR: um termo por peça/concretagem, preenchido na VÉSPERA a partir do Checklist Mestre. "
 "Anexar o checklist e as fotos. Alterações de armação/fôrma/escoramento são alçada do projetista (NBR 6118:2023/15696).")
rn.italic=True; rn.font.size=Pt(8.5); rn.font.color.rgb=REALCE

h("1. IDENTIFICAÇÃO")
for lbl,hint in [("Obra:","nome / endereço"),("Peça / pavimento:","ex.: laje L2 + vigas"),
 ("Data da inspeção (véspera):","dd/mm/aaaa"),("Concretagem prevista:","dd/mm — h"),
 ("Projeto (revisão) usado:","estrutural rev. X"),("Responsável pela inspeção:","nome / CREA")]:
    campo(lbl,hint)

h("2. SÍNTESE DA INSPEÇÃO (do Checklist Mestre anexo)")
cols=["Sistema","Status","Observações"]
t=doc.add_table(rows=1,cols=3); t.style="Table Grid"
for i,c in enumerate(cols):
    cell=t.rows[0].cells[i]; shade(cell,"0E3A34")
    rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xF5,0xEF,0xE0); rr.font.size=Pt(9)
for sist in ["Fôrma","Armação (incl. cobrimento/espaçadores)","Escoramento","Embutidos (elétrica/hidráulica/esperas)","Limpeza e acesso"]:
    cells=t.add_row().cells
    cells[0].paragraphs[0].add_run(sist).font.size=Pt(9)
    for k in (1,2):
        rr=cells[k].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(9)

h("3. PENDÊNCIAS (se houver — obrigatório prazo + reconferência)")
tabela(["#","Pendência","Responsável","Prazo","Reconferido em (h/rubrica)"],3,"1C5A4E")

h("4. VEREDITO")
p=doc.add_paragraph()
p.add_run("(  ) LIBERADO    (  ) LIBERADO COM PENDÊNCIAS (item 3)    (  ) TRAVADO — motivo: ").bold=True
rr=p.add_run("[descrever com fato + solução + novo cronograma]"); rr.font.color.rgb=REALCE; rr.italic=True

h("5. CONFERÊNCIA DOS EMBUTIDOS (assinatura das equipes)")
tabela(["Equipe","Nome","Assinatura"],2,"1C5A4E")

h("6. RECONFERÊNCIA NO DIA (15 min antes do caminhão)")
campo("Hora / itens reconferidos / rubrica:","…")

doc.add_paragraph()
ass=doc.add_paragraph(); ass.alignment=WD_ALIGN_PARAGRAPH.CENTER
ass.add_run("_______________________________________________\n").font.color.rgb=CINZA
ra=ass.add_run("Responsável pela inspeção — assinatura / CREA\n"); ra.bold=True; ra.font.color.rgb=VERDE
foot=doc.add_paragraph(); foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
rf=foot.add_run("Modelo revisado e assinado por Gustavo Domingos — Eng. Civil, CREA-PR 140.964-D · Escola de Obra · engenhariagd.com.br")
rf.font.size=Pt(8); rf.font.color.rgb=CINZA

out=os.path.abspath(os.path.join(os.path.dirname(__file__),"..","06-curso-inspecao","kit","relatorio","termo-liberacao.docx"))
doc.save(out); print("TERMO DOCX GERADO:",out,f"({os.path.getsize(out)} bytes)")
