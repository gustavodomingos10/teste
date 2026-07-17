#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Parecer/Relatório de Fundação (blindado) em .docx — Kit do Curso 2.
Triagem e acompanhamento à luz da NBR 6122:2019; a decisão e o dimensionamento são do projetista."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERDE=RGBColor(0x0E,0x3A,0x34); DOUR=RGBColor(0xC9,0xA2,0x4B); CINZA=RGBColor(0x6B,0x6B,0x63); REALCE=RGBColor(0x9C,0x65,0x00)
doc=Document(); st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexc); tcPr.append(sh)
def h(txt,size=13,sb=10):
    p=doc.add_paragraph(); r=p.add_run(txt); r.bold=True; r.font.size=Pt(size); r.font.color.rgb=VERDE; r.font.name="Georgia"
    p.paragraph_format.space_before=Pt(sb); p.paragraph_format.space_after=Pt(3); return p
def campo(label,hint):
    p=doc.add_paragraph(); r=p.add_run(label+" "); r.bold=True; r.font.color.rgb=VERDE
    r2=p.add_run("["+hint+"]"); r2.font.color.rgb=REALCE; r2.italic=True; return p

tt=doc.add_paragraph(); tt.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tt.add_run("PARECER TÉCNICO — TRIAGEM E ACOMPANHAMENTO DE FUNDAÇÃO"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=VERDE; r.font.name="Georgia"
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
rs=sub.add_run("Modelo Blindado · Escola de Obra · GD Engenharia e Perícia"); rs.font.color.rgb=DOUR; rs.bold=True; rs.font.size=Pt(9)
sel=doc.add_paragraph(); sel.alignment=WD_ALIGN_PARAGRAPH.CENTER
rsel=sel.add_run("Padrão Diamante — o rigor de quem assina"); rsel.italic=True; rsel.font.color.rgb=CINZA; rsel.font.size=Pt(8.5)

nota=doc.add_paragraph(); rn=nota.add_run(
 "COMO USAR: substitua o texto entre [colchetes]. Este parecer documenta a LEITURA, a TRIAGEM e o "
 "ACOMPANHAMENTO da execução; a definição e o dimensionamento da fundação são do projetista de "
 "fundações/geotécnico (NBR 6122:2019). Apague esta nota ao finalizar.")
rn.italic=True; rn.font.size=Pt(8.5); rn.font.color.rgb=REALCE

h("1. IDENTIFICAÇÃO")
for lbl,hint in [("Solicitante:","nome / contato"),("Obra / Endereço:","logradouro, nº, cidade/UF"),
 ("Tipo de obra / uso:","residencial, edifício, etc."),("Data:","dd/mm/aaaa"),
 ("Projetista de fundações:","nome / CREA (responsável pelo projeto)"),
 ("Responsável por este parecer:","Gustavo Domingos — Eng. Civil, CREA-PR 140.964-D")]:
    campo(lbl,hint)

h("2. OBJETO, ESCOPO E LIMITAÇÕES")
doc.add_paragraph("Objeto: triagem da família de fundação e acompanhamento da execução conforme o projeto de fundações.")
campo("Escopo:","o que foi solicitado/acompanhado")
p=doc.add_paragraph(); p.add_run("Limitações: ").bold=True
doc.add_paragraph("Este parecer NÃO constitui projeto de fundações. A escolha do tipo, o dimensionamento e a "
 "responsabilidade técnica pelo projeto são do projetista de fundações/geotécnico (NBR 6122:2019). As hipóteses de "
 "triagem baseiam-se na sondagem e nas cargas informadas.")

h("3. DADOS DE PARTIDA")
campo("Sondagem (nº / data / responsável):","SP-XX, dd/mm/aaaa — NBR 6484:2020")
campo("Camada resistente (profundidade):","ex.: N-SPT firma a X m")
campo("Nível d'água (NA):","ex.: a Y m / ausente")
campo("Cargas informadas (projeto estrutural):","ordem de grandeza por apoio")
campo("Restrições da obra:","vizinhança, acesso, prazo/custo")

h("4. TRIAGEM (leitura à luz da NBR 6122:2019)")
doc.add_paragraph("Registro do raciocínio de triagem. Família provável e o porquê — sem cravar tipo/dimensionamento.")
campo("Família provável:","RASA / PROFUNDA / a definir (zona cinzenta)")
campo("Justificativa:","camada resistente, NA, carga e restrições")
campo("Tipos candidatos (a validar com o projetista):","ex.: hélice contínua / escavada (vizinhança sensível)")

h("5. DEFINIÇÃO DO PROJETISTA (recebida)")
cols=["Item","Definido pelo projetista"]
t=doc.add_table(rows=1,cols=len(cols)); t.style="Table Grid"
for i,c in enumerate(cols):
    cell=t.rows[0].cells[i]; shade(cell,"0E3A34")
    rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xF5,0xEF,0xE0); rr.font.size=Pt(9)
for item in ["Tipo de fundação","Geometria / dimensões","Cotas / comprimentos","Observações do projeto"]:
    cells=t.add_row().cells
    cells[0].paragraphs[0].add_run(item).font.size=Pt(9)
    rr=cells[1].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(9)

h("6. ACOMPANHAMENTO DA EXECUÇÃO")
cols2=["#","Verificação","Conforme?","Observação"]
t2=doc.add_table(rows=1,cols=len(cols2)); t2.style="Table Grid"
for i,c in enumerate(cols2):
    cell=t2.rows[0].cells[i]; shade(cell,"1C5A4E")
    rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); rr.font.size=Pt(9)
for it in ["Locação e cota de assentamento","Dimensões / armação (conforme projeto)","Controle por elemento (estaca: comprimento/nega; sapata: apoio)","Provas de carga (quando exigidas) — NBR 6489/12131/16903","Registro fotográfico datado"]:
    cells=t2.add_row().cells
    cells[0].paragraphs[0].add_run("[ ]").font.size=Pt(9)
    cells[1].paragraphs[0].add_run(it).font.size=Pt(9)
    for k in (2,3):
        rr=cells[k].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(9)

h("7. DIVERGÊNCIAS E ENCAMINHAMENTOS")
campo("Divergências observadas:","e a quem foram comunicadas, com data")
campo("Encaminhamentos (prazo / responsável):","ação, prazo e responsável")

h("8. CONCLUSÃO")
campo("Conclusão:","síntese da triagem e do acompanhamento; sem promessa de resultado")

h("9. RESPONSABILIDADE TÉCNICA")
doc.add_paragraph("A definição e o dimensionamento da fundação são de responsabilidade do projetista de fundações. "
 "Este parecer refere-se à leitura, triagem e acompanhamento da execução conforme projeto.")
doc.add_paragraph()
ass=doc.add_paragraph(); ass.alignment=WD_ALIGN_PARAGRAPH.CENTER
ass.add_run("_______________________________________________\n").font.color.rgb=CINZA
ra=ass.add_run("Gustavo Domingos — Engenheiro Civil\nCREA-PR 140.964-D · CREA-SP 5071652757\n"); ra.bold=True; ra.font.color.rgb=VERDE
ass.add_run("Perito judicial atuante em processos do TJ-PR e do TJ-SP  [VALIDAR redação — D1]").font.color.rgb=CINZA
foot=doc.add_paragraph(); foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
rf=foot.add_run("Escola de Obra · GD Engenharia e Perícia · engenhariagd.com.br"); rf.font.size=Pt(8); rf.font.color.rgb=CINZA

out=os.path.abspath(os.path.join(os.path.dirname(__file__),"..","04-curso-fundacoes","kit","relatorio","parecer-fundacao.docx"))
doc.save(out); print("PARECER DOCX GERADO:",out,f"({os.path.getsize(out)} bytes)")
