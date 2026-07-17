#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilha de Controle de Concreto (Kit do Curso 3).
Abas: Parametros, Recebimento (uma linha por caminhão, veredito automático),
CPs (controle de corpos de prova), Dashboard.
Saída: 05-curso-concreto/kit/planilha/planilha-concreto.xlsx"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

VERDE="0E3A34"; VMED="1C5A4E"; CREME="F5EFE0"; CREME2="FBF8F0"; DOUR="C9A24B"; CARVAO="20211E"; CINZA="6B6B63"
f_tit=Font(name="Georgia",size=15,bold=True,color=CREME)
f_hdr=Font(name="Calibri",size=10,bold=True,color=CREME)
f_lbl=Font(name="Calibri",size=10,bold=True,color=VERDE)
f_norm=Font(name="Calibri",size=10,color=CARVAO)
f_small=Font(name="Calibri",size=8.5,italic=True,color=CINZA)
fill_v=PatternFill("solid",fgColor=VERDE); fill_vm=PatternFill("solid",fgColor=VMED)
fill_c2=PatternFill("solid",fgColor=CREME2); fill_d=PatternFill("solid",fgColor=DOUR)
fill_ok=PatternFill("solid",fgColor="D9EAD3"); fill_warn=PatternFill("solid",fgColor="FFF2CC"); fill_bad=PatternFill("solid",fgColor="F4CCCC")
thin=Side(style="thin",color="D9D2BE"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal="center",vertical="center",wrap_text=True)
left=Alignment(horizontal="left",vertical="center",wrap_text=True)

wb=Workbook()

# ---------- PARAMETROS ----------
wp=wb.active; wp.title="Parametros"; wp.sheet_view.showGridLines=False
wp["A1"]="PARÂMETROS DO CONTROLE — ajuste conforme o pedido e assine"
wp["A1"].font=f_tit; wp["A1"].fill=fill_v; wp.merge_cells("A1:D1"); wp.row_dimensions[1].height=26
wp["A1"].alignment=Alignment(horizontal="left",vertical="center")
params=[
 ("Slump pedido (mm)",100,"Conforme pedido/projeto. Confira na nota fiscal."),
 ("Tolerância de slump (± mm)",20,"Faixa da NBR 7212 conforme o slump pedido. [VALIDAR F4]"),
 ("Tempo-limite descarga (min)",150,"Da adição da água ao fim da descarga (NBR 7212; aditivos podem alterar). [VALIDAR F2]"),
 ("fck de projeto (MPa)",30,"Conforme projeto estrutural (classe NBR 8953)."),
]
wp["A3"]="Parâmetro"; wp["B3"]="Valor"; wp["C3"]="Observação / fonte"
for c in ("A3","B3","C3"): wp[c].font=f_hdr; wp[c].fill=fill_vm; wp[c].alignment=left; wp[c].border=border
r=4
for nome,val,obs in params:
    wp[f"A{r}"]=nome; wp[f"A{r}"].font=f_lbl
    wp[f"B{r}"]=val; wp[f"B{r}"].font=f_norm; wp[f"B{r}"].fill=fill_d; wp[f"B{r}"].alignment=center
    wp[f"C{r}"]=obs; wp[f"C{r}"].font=f_small; wp[f"C{r}"].alignment=left
    for col in "ABC": wp[f"{col}{r}"].border=border
    r+=1
wp["A9"]="⚠ A aceitação definitiva do concreto segue os critérios estatísticos da NBR 12655 — a coluna 'situação' da aba CPs é uma triagem de alerta, não o critério final. [VALIDAR F7]"
wp["A9"].font=f_small; wp.merge_cells("A9:D10"); wp["A9"].alignment=left
wp.column_dimensions["A"].width=30; wp.column_dimensions["B"].width=12; wp.column_dimensions["C"].width=48

# ---------- RECEBIMENTO ----------
ws=wb.create_sheet("Recebimento"); ws.sheet_view.showGridLines=False
ws["A1"]="RECEBIMENTO — uma linha por caminhão (aceitação provisória, NBR 12655/7212/16889)"
ws["A1"].font=f_tit; ws["A1"].fill=fill_v; ws.merge_cells("A1:L1"); ws.row_dimensions[1].height=24
ws["A1"].alignment=Alignment(horizontal="left",vertical="center")
headers=[("A","Data",11),("B","NF nº",10),("C","Hora água (central)",11),("D","Fim descarga",11),
 ("E","Tempo decorrido (min)",11),("F","Slump medido (mm)",11),("G","Nota confere?",10),
 ("H","Aspecto normal?",10),("I","Água adicionada?",10),("J","Veredito",16),
 ("K","CPs moldados (ids)",16),("L","Observações",24)]
for col,txt,w in headers:
    c=ws[f"{col}2"]; c.value=txt; c.font=f_hdr; c.fill=fill_v; c.alignment=center; c.border=border
    ws.column_dimensions[col].width=w
ws.row_dimensions[2].height=30; ws.freeze_panes="A3"
FIRST,LAST=3,102
for row in range(FIRST,LAST+1):
    ws[f"E{row}"]=f'=IF(OR(C{row}="",D{row}=""),"",ROUND((D{row}-C{row})*1440,0))'
    ws[f"J{row}"]=(f'=IF(F{row}="","",'
      f'IF(OR(G{row}="Nao",H{row}="Nao",I{row}="Sim"),"RECUSAR/REGISTRAR",'
      f'IF(AND(E{row}<>"",E{row}>Parametros!$B$6),"RECUSAR - tempo",'
      f'IF(ABS(F{row}-Parametros!$B$4)<=Parametros!$B$5,"LIBERADO",'
      f'IF(F{row}<Parametros!$B$4,"CORRIGIR VIA CENTRAL","RECUSAR - fluido")))))')
    for col,_,_ in headers:
        cell=ws[f"{col}{row}"]; cell.font=f_norm; cell.border=border
        cell.alignment=center if col in "ABCDEFGHIJ" else left
        if row%2==0: cell.fill=fill_c2
    ws[f"C{row}"].number_format="hh:mm"; ws[f"D{row}"].number_format="hh:mm"
    ws[f"A{row}"].number_format="dd/mm/yyyy"
for col in ("G","H"):
    dv=DataValidation(type="list",formula1='"Sim,Nao"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}{FIRST}:{col}{LAST}")
dv=DataValidation(type="list",formula1='"Sim,Nao"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"I{FIRST}:I{LAST}")
rng=f"J{FIRST}:J{LAST}"
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"LIBERADO"'],fill=fill_ok,font=Font(color="1b4d1f",bold=True)))
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"CORRIGIR VIA CENTRAL"'],fill=fill_warn,font=Font(color="9C6500",bold=True)))
ws.conditional_formatting.add(rng,CellIsRule(operator="beginsWith",formula=['"RECUSAR"'],fill=fill_bad,font=Font(color="7f1d1d",bold=True)))

# ---------- CPs ----------
wc=wb.create_sheet("CPs"); wc.sheet_view.showGridLines=False
wc["A1"]="CORPOS DE PROVA — moldagem NBR 5738 · ensaio NBR 5739 · aceitação NBR 12655"
wc["A1"].font=f_tit; wc["A1"].fill=fill_v; wc.merge_cells("A1:H1"); wc.row_dimensions[1].height=24
wc["A1"].alignment=Alignment(horizontal="left",vertical="center")
h2=[("A","CP id",10),("B","NF / caminhão",12),("C","Data moldagem",12),("D","Peça concretada",18),
 ("E","Idade (dias)",9),("F","Resultado (MPa)",11),("G","Situação (triagem)",18),("H","Observações",24)]
for col,txt,w in h2:
    c=wc[f"{col}2"]; c.value=txt; c.font=f_hdr; c.fill=fill_v; c.alignment=center; c.border=border
    wc.column_dimensions[col].width=w
wc.freeze_panes="A3"
for row in range(3,103):
    wc[f"G{row}"]=(f'=IF(F{row}="","",'
      f'IF(F{row}>=Parametros!$B$7,"OK (>= fck)",'
      f'"ALERTA: abaixo do fck - aplicar criterios NBR 12655 / acionar RT"))')
    for col,_,_ in h2:
        cell=wc[f"{col}{row}"]; cell.font=f_norm; cell.border=border
        cell.alignment=center if col in "ABCEFG" else left
        if row%2==0: cell.fill=fill_c2
    wc[f"C{row}"].number_format="dd/mm/yyyy"; wc[f"F{row}"].number_format="0.0"
rng2="G3:G102"
wc.conditional_formatting.add(rng2,CellIsRule(operator="beginsWith",formula=['"OK"'],fill=fill_ok,font=Font(color="1b4d1f",bold=True)))
wc.conditional_formatting.add(rng2,CellIsRule(operator="beginsWith",formula=['"ALERTA"'],fill=fill_bad,font=Font(color="7f1d1d",bold=True)))

# ---------- DASHBOARD ----------
wd=wb.create_sheet("Dashboard"); wd.sheet_view.showGridLines=False
wd["A1"]="DASHBOARD — Controle do Concreto"; wd["A1"].font=f_tit; wd["A1"].fill=fill_v
wd.merge_cells("A1:E1"); wd.row_dimensions[1].height=26; wd["A1"].alignment=Alignment(horizontal="left",vertical="center")
def metric(lc,vc,label,formula,fill):
    wd[lc]=label; wd[lc].font=f_lbl; wd[lc].border=border; wd[lc].alignment=left
    wd[vc]=formula; wd[vc].font=Font(name="Georgia",size=18,bold=True,color=VERDE); wd[vc].alignment=center; wd[vc].fill=fill; wd[vc].border=border
metric("A3","A4","Caminhões",'=COUNTA(Recebimento!B3:B102)',fill_c2)
metric("B3","B4","Liberados",'=COUNTIF(Recebimento!J3:J102,"LIBERADO")',fill_ok)
metric("C3","C4","Corrigidos",'=COUNTIF(Recebimento!J3:J102,"CORRIGIR VIA CENTRAL")',fill_warn)
metric("D3","D4","Recusados",'=COUNTIF(Recebimento!J3:J102,"RECUSAR*")',fill_bad)
metric("A6","A7","CPs registrados",'=COUNTA(CPs!A3:A102)',fill_c2)
metric("B6","B7","CPs OK",'=COUNTIF(CPs!G3:G102,"OK*")',fill_ok)
metric("C6","C7","CPs em alerta",'=COUNTIF(CPs!G3:G102,"ALERTA*")',fill_bad)
wd["A9"]=('Como usar: uma linha por caminhão no Recebimento (veredito automático a partir dos Parametros); '
 'uma linha por CP na aba CPs. A "situação" dos CPs é triagem de alerta — a aceitação definitiva segue os '
 'critérios estatísticos da NBR 12655, aplicados pelo responsável técnico.')
wd["A9"].font=f_small; wd.merge_cells("A9:E11"); wd["A9"].alignment=left
wd["A13"]="Revisado e assinado por Gustavo Domingos — Engenheiro Civil, CREA-PR 140.964-D"
wd["A13"].font=Font(name="Calibri",size=9,bold=True,color=DOUR)
for w,c in [(16,"A"),(12,"B"),(12,"C"),(12,"D"),(10,"E")]: wd.column_dimensions[c].width=w

out=os.path.abspath(os.path.join(os.path.dirname(__file__),"..","05-curso-concreto","kit","planilha","planilha-concreto.xlsx"))
wb.save(out); print("PLANILHA GERADA:",out)
