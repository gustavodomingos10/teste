#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Testa a planilha: injeta 3 casos, avalia as fórmulas de verdade (lib 'formulas')
e confere Situação, Gravidade, Encaminhamento e o Dashboard."""
import os, re, sys, shutil, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
import formulas

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SRC = os.path.join(BASE, "02-curso-piloto", "kit", "planilha", "planilha-fissuras.xlsx")
TMP = "/tmp/planilha-teste.xlsx"
shutil.copy(SRC, TMP)

# --- Injeta 3 casos (colunas: A Nº,B amb,C orient,D local,E ini,F d1,G atual,H d2,K origem,L umid,M risco) ---
casos = [
    # Nº, ambiente, orient, local, ini, atual, origem, umidade, risco, ESPERADO(sit,grav,enc)
    (1,"Sala - parede","Diagonal 45","Canto de abertura",0.25,0.28,"Concentracao de tensao (canto)","Nao","Nao",
        ("ESTAVEL","BAIXA","Monitorar e registrar")),
    (2,"Fachada - reboco","Mapeada","Topo da parede",0.20,0.45,"Retracao","Nao","Nao",
        ("ATIVA","MEDIA","Tratar a causa + monitorar")),
    (3,"Garagem - viga/parede","Diagonal 45","Muro",0.50,0.90,"Recalque diferencial","Nao","Nao",
        ("ATIVA","ALTA","Escorar/interditar e acionar calculista")),
]
wb = load_workbook(TMP)
ws = wb["Lancamentos"]
for i,(n,amb,ori,loc,ini,atu,org,umi,ris,_) in enumerate(casos):
    r = 3+i
    ws[f"A{r}"]=n; ws[f"B{r}"]=amb; ws[f"C{r}"]=ori; ws[f"D{r}"]=loc
    ws[f"E{r}"]=ini; ws[f"G{r}"]=atu; ws[f"K{r}"]=org; ws[f"L{r}"]=umi; ws[f"M{r}"]=ris
wb.save(TMP)

# --- Avalia com a lib formulas ---
xl = formulas.ExcelModel().loads(TMP).finish()
sol = xl.calculate()
def get(sheet, cell):
    pat = re.compile(r"\]"+sheet.upper()+r"'?\!"+cell.upper()+r"$")
    for k,v in sol.items():
        if pat.search(k):
            try: val = v.value[0,0]
            except Exception:
                try: val = v.value
                except Exception: val = v
            if hasattr(val,"ravel"):
                val = val.ravel()[0]
            return val
    return None

ok = True
print("="*66)
print(" TESTE DA PLANILHA — 3 casos avaliados com engine de fórmulas")
print("="*66)
for i,(n,amb,ori,loc,ini,atu,org,umi,ris,esp) in enumerate(casos):
    r = 3+i
    sit = str(get("Lancamentos", f"J{r}")).strip().strip('"')
    grav = str(get("Lancamentos", f"N{r}")).strip().strip('"')
    enc = str(get("Lancamentos", f"O{r}")).strip().strip('"')
    e_sit,e_grav,e_enc = esp
    passed = (sit==e_sit and grav==e_grav and enc==e_enc)
    ok = ok and passed
    print(f"\nCaso {n}: {amb}  (ini={ini} atual={atu} Δ={round(atu-ini,2)} origem={org} risco={ris})")
    print(f"  Situação:      obtido={sit:8}  esperado={e_sit:8}  {'✔' if sit==e_sit else '�’✗'}")
    print(f"  Gravidade:     obtido={grav:8}  esperado={e_grav:8}  {'✔' if grav==e_grav else '✗'}")
    print(f"  Encaminhamento: {enc}")
    print(f"     esperado:     {e_enc}  {'✔' if enc==e_enc else '✗'}")
    print(f"  => {'PASSOU' if passed else 'FALHOU'}")

# --- Dashboard ---
tot = get("Dashboard","A4"); ativ = get("Dashboard","B4"); est = get("Dashboard","C4")
ga = get("Dashboard","B8"); gm = get("Dashboard","B9"); gb = get("Dashboard","B10")
def num(x):
    try: return int(round(float(x)))
    except Exception: return x
print("\n" + "-"*66)
print(" DASHBOARD")
print(f"  Total fissuras: {num(tot)} (esperado 3)  {'✔' if num(tot)==3 else '✗'}")
print(f"  Ativas: {num(ativ)} (esperado 2)  {'✔' if num(ativ)==2 else '✗'}")
print(f"  Estáveis: {num(est)} (esperado 1)  {'✔' if num(est)==1 else '✗'}")
print(f"  Gravidade ALTA/MEDIA/BAIXA: {num(ga)}/{num(gm)}/{num(gb)} (esperado 1/1/1)  "
      f"{'✔' if (num(ga),num(gm),num(gb))==(1,1,1) else '✗'}")
dash_ok = (num(tot)==3 and num(ativ)==2 and num(est)==1 and (num(ga),num(gm),num(gb))==(1,1,1))
ok = ok and dash_ok
print("\n" + "="*66)
print(" RESULTADO FINAL:", "TODOS OS TESTES PASSARAM ✔" if ok else "HÁ FALHAS ✗")
print("="*66)
sys.exit(0 if ok else 1)
