#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Registro de Concretagem Blindado (.docx) — Kit do Curso 3."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERDE=RGBColor(0x0E,0x3A,0x34); DOUR=RGBColor(0xC9,0xA2,0x4B); CINZA=RGBColor(0x6B,0x6B,0x63); REALCE=RGBColor(0x9C,0x65,0x00)
doc=Document(); st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(10.5)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexc); tcPr.append(sh)
def h(txt,size=13,sb=10):
    p=doc.add_paragraph(); r=p.add_run(txt); r.bold=True; r.font.size=Pt(size); r.font.color.rgb=VERDE; r.font.name="Georgia"
    p.paragraph_format.space_before=Pt(sb); p.paragraph_format.space_after=Pt(3); return p
def campo(label,hint):
    p=doc.add_paragraph(); r=p.add_run(label+" "); r.bold=True; r.font.color.rgb=VERDE
    r2=p.add_run("["+hint+"]"); r2.font.color.rgb=REALCE; r2.italic=True; return p
def tabela(cols, rows, fillhdr="0E3A34"):
    t=doc.add_table(rows=1,cols=len(cols)); t.style="Table Grid"
    for i,c in enumerate(cols):
        cell=t.rows[0].cells[i]; shade(cell,fillhdr)
        rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xF5,0xEF,0xE0); rr.font.size=Pt(9)
    for _ in range(rows):
        cells=t.add_row().cells
        for i in range(len(cols)):
            rr=cells[i].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(9)
    return t

tt=doc.add_paragraph(); tt.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tt.add_run("REGISTRO DE CONCRETAGEM — DOSSIÊ BLINDADO"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=VERDE; r.font.name="Georgia"
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
rs=sub.add_run("Escola de Obra · GD Engenharia e Perícia · Padrão Diamante"); rs.font.color.rgb=DOUR; rs.bold=True; rs.font.size=Pt(9)
nota=doc.add_paragraph()
rn=nota.add_run("COMO USAR: um registro por concretagem (a planilha consolida a obra). Substitua os [colchetes]. "
 "Anexe: notas fiscais, fotos dos ensaios e certificados de rompimento. Apague esta nota ao finalizar.")
rn.italic=True; rn.font.size=Pt(8.5); rn.font.color.rgb=REALCE

h("1. IDENTIFICAÇÃO DA CONCRETAGEM")
for lbl,hint in [("Obra:","nome / endereço"),("Data:","dd/mm/aaaa"),("Peça(s) concretada(s):","ex.: laje L2, vigas V5-V8"),
 ("Volume total (m³):","…"),("Central / fornecedor:","…"),
 ("fck de projeto / classe:","ex.: 30 MPa / C30 (NBR 8953)"),("Slump pedido (mm) e tolerância:","ex.: 100 ± [VALIDAR F4]"),
 ("Responsável pelo recebimento:","nome / CREA")]:
    campo(lbl,hint)

h("2. CAMINHÕES RECEBIDOS (um por linha — anexar notas)")
tabela(["NF nº","Hora água","Chegada","Início/Fim descarga","Slump (mm)","Veredito","CPs (ids)"],4)
p=doc.add_paragraph(); rr=p.add_run("Veredito: LIBERADO · CORRIGIDO VIA CENTRAL (re-ensaio anexo) · RECUSADO (motivo no item 4). "
 "Tempo-limite conforme NBR 7212 [VALIDAR F2]."); rr.font.size=Pt(8.5); rr.font.color.rgb=CINZA

h("3. ENSAIOS DE ABATIMENTO (NBR 16889)")
tabela(["NF nº","Hora","Resultado (mm)","Dentro da tolerância?","Foto (ref.)","Executado por"],4)

h("4. OCORRÊNCIAS E RECUSAS")
campo("Ocorrências (água solicitada/adicionada, correções via central, recusas com motivo/hora):","descrever com hora e testemunhas")

h("5. CORPOS DE PROVA (NBR 5738)")
tabela(["CP id","NF nº","Peça","Data moldagem","Cura inicial ok?","Envio ao laboratório"],4)

h("6. CURA DA PEÇA")
campo("Método e período de cura:","ex.: aspersão/manta por [X] dias — NBR 14931")
campo("Responsável pela cura:","nome")

h("7. RESULTADOS E ACEITAÇÃO DEFINITIVA (NBR 5739 / 12655)")
tabela(["CP id","Idade (dias)","Resultado (MPa)","Situação","Certificado (ref.)"],4)
p=doc.add_paragraph(); rr=p.add_run("Resultado abaixo do esperado: aplicar os critérios da NBR 12655; ensaios complementares "
 "(extração NBR 7680, esclerometria) e decisão do responsável técnico/projetista. [VALIDAR F10]"); rr.font.size=Pt(8.5); rr.font.color.rgb=CINZA

h("8. RESPONSABILIDADE")
doc.add_paragraph("Este registro documenta a aceitação provisória do concreto fresco e o controle de aceitação. "
 "Decisões estruturais decorrentes de resultados não conformes cabem ao responsável técnico/projetista estrutural.")
doc.add_paragraph()
ass=doc.add_paragraph(); ass.alignment=WD_ALIGN_PARAGRAPH.CENTER
ass.add_run("_______________________________________________\n").font.color.rgb=CINZA
ra=ass.add_run("Responsável pelo recebimento — assinatura / CREA\n"); ra.bold=True; ra.font.color.rgb=VERDE
foot=doc.add_paragraph(); foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
rf=foot.add_run("Modelo revisado e assinado por Gustavo Domingos — Eng. Civil, CREA-PR 140.964-D · Escola de Obra · engenhariagd.com.br")
rf.font.size=Pt(8); rf.font.color.rgb=CINZA

out=os.path.abspath(os.path.join(os.path.dirname(__file__),"..","05-curso-concreto","kit","relatorio","registro-concretagem.docx"))
doc.save(out); print("REGISTRO DOCX GERADO:",out,f"({os.path.getsize(out)} bytes)")
