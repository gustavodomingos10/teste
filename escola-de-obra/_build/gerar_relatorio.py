#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Modelo de Relatório Técnico Blindado (.docx) — Kit Diamante.
Estrutura NBR 13752 / NBR 16747: identificação, escopo/limitações, metodologia,
registro fotográfico padronizado (2 colunas), análise causa raiz, recomendações
com prazo/responsável, encaminhamentos, assinatura/CREA.
Campos preenchíveis destacados em [colchetes] realçados."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERDE = RGBColor(0x0E,0x3A,0x34); DOUR = RGBColor(0xC9,0xA2,0x4B); CINZA = RGBColor(0x6B,0x6B,0x63)
REALCE = RGBColor(0x9C,0x65,0x00)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexcolor)
    tcPr.append(sh)

def h(txt, size=13, color=VERDE, space_before=10):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    r.font.name = "Georgia"; p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(3)
    return p

def campo(label, hint):
    p = doc.add_paragraph()
    r = p.add_run(label + " "); r.bold = True; r.font.color.rgb = VERDE
    r2 = p.add_run("[" + hint + "]"); r2.font.color.rgb = REALCE; r2.italic = True
    return p

# ---------- Cabeçalho ----------
tt = doc.add_paragraph(); tt.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tt.add_run("RELATÓRIO TÉCNICO — DIAGNÓSTICO DE FISSURAS E TRINCAS")
r.bold = True; r.font.size = Pt(15); r.font.color.rgb = VERDE; r.font.name = "Georgia"
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Modelo Blindado · Escola de Obra · GD Engenharia e Perícia"); rs.font.color.rgb = DOUR; rs.bold = True; rs.font.size = Pt(9)
sel = doc.add_paragraph(); sel.alignment = WD_ALIGN_PARAGRAPH.CENTER
rsel = sel.add_run("Padrão Diamante — o rigor de quem assina"); rsel.italic = True; rsel.font.color.rgb = CINZA; rsel.font.size = Pt(8.5)

nota = doc.add_paragraph()
rn = nota.add_run("COMO USAR: substitua todo texto entre [colchetes] pelos dados do caso. "
                  "Não apague as seções — a estrutura é a sua blindagem jurídica. Apague esta nota ao finalizar.")
rn.italic = True; rn.font.size = Pt(8.5); rn.font.color.rgb = REALCE

# 1. Identificação
h("1. IDENTIFICAÇÃO")
for lbl,hint in [("Solicitante:","nome / contato"),("Endereço do imóvel:","logradouro, nº, cidade/UF"),
                 ("Tipo de imóvel / uso:","residencial, sobrado, etc."),("Data da vistoria:","dd/mm/aaaa"),
                 ("Nº do caso / OS:","identificador interno"),("Responsável técnico:","Gustavo Domingos — Eng. Civil, CREA-PR 140.964-D")]:
    campo(lbl,hint)

# 2. Escopo e limitações
h("2. OBJETO, ESCOPO E LIMITAÇÕES")
doc.add_paragraph("Objeto: diagnóstico do quadro fissuratório observado, quanto à origem provável, atividade e gravidade, "
                  "com recomendação de encaminhamento.")
campo("Escopo contratado:","descrever o que foi solicitado")
p = doc.add_paragraph(); r = p.add_run("Limitações: ").bold = True
doc.add_paragraph("Vistoria de caráter [visual/cautelar], sem ensaios destrutivos, salvo indicação em contrário. "
                  "As hipóteses de origem são baseadas em inspeção visual e medição de abertura; a confirmação de "
                  "causas estruturais ou de fundação pode exigir ensaios/investigações específicas, indicados nas recomendações. "
                  "Este relatório reflete a condição observada na data da vistoria.", style=None)

# 3. Metodologia
h("3. METODOLOGIA")
for item in ["Inspeção do geral ao particular (contexto → detalhe), conforme boa prática de inspeção predial (NBR 16747).",
             "Caracterização de cada fissura: orientação, localização, abertura medida com fissurômetro e registro fotográfico com escala.",
             "Avaliação de atividade por monitoramento (selo de gesso/comparador) quando aplicável.",
             "Enquadramento de gravidade e encaminhamento conforme critério técnico (planilha e fluxograma do método)."]:
    doc.add_paragraph(item, style="List Bullet")
p = doc.add_paragraph(); rr = p.add_run("Normas de referência: "); rr.bold=True; rr.font.color.rgb=VERDE
p.add_run("NBR 6118 (estruturas de concreto), NBR 6122 (fundações), NBR 16747 (inspeção predial), "
          "NBR 13752 (perícias de engenharia). [Confirmar edição vigente na data.]")

# 4. Registro fotográfico (2 colunas)
h("4. REGISTRO FOTOGRÁFICO PADRONIZADO")
doc.add_paragraph("Cada fissura numerada corresponde ao mapa fissuratório (item 5). Legenda: nº, local, orientação, abertura.")
tbl = doc.add_table(rows=2, cols=2); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in tbl.rows:
    for cell in row.cells:
        cell.width = Cm(8)
        ph = cell.paragraphs[0]; ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rp = ph.add_run("[INSERIR FOTO]\n[Foto nº __ — local / orientação / abertura mm]")
        rp.font.size = Pt(8.5); rp.font.color.rgb = REALCE; rp.italic = True
doc.add_paragraph("(Duplique a tabela conforme o número de fissuras registradas.)", style=None).runs[0].italic = True

# 5. Mapa fissuratório e caracterização
h("5. MAPA FISSURATÓRIO E CARACTERIZAÇÃO")
cols = ["Nº","Local/Elemento","Orient.","Abertura (mm)","Origem (hipótese)","Situação","Gravidade"]
t2 = doc.add_table(rows=1, cols=len(cols)); t2.style = "Table Grid"
for i,c in enumerate(cols):
    cell = t2.rows[0].cells[i]; shade(cell,"0E3A34")
    rr = cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xF5,0xEF,0xE0); rr.font.size=Pt(8.5)
for _ in range(3):
    cells = t2.add_row().cells
    for i in range(len(cols)):
        rr = cells[i].paragraphs[0].add_run("[ ]"); rr.font.color.rgb=REALCE; rr.font.size=Pt(8.5)

# 6. Análise / causa raiz
h("6. ANÁLISE E CAUSA RAIZ")
doc.add_paragraph("Descreva o padrão dominante do quadro fissuratório e a causa raiz mais provável, "
                  "distinguindo sintoma de causa. Fundamente com os achados dos itens 4 e 5.")
campo("Padrão dominante observado:","ex.: fissuras diagonais concordantes apontando para o canto X")
campo("Causa raiz provável:","ex.: recalque diferencial de fundação a investigar (NBR 6122)")
campo("Causas secundárias / associadas:","ex.: umidade ascendente na base (NBR 9575)")

# 7. Recomendações com prazo e responsável
h("7. RECOMENDAÇÕES (com prazo e responsável)")
cols2 = ["#","Recomendação","Prazo","Responsável","Prioridade"]
t3 = doc.add_table(rows=1, cols=len(cols2)); t3.style="Table Grid"
for i,c in enumerate(cols2):
    cell=t3.rows[0].cells[i]; shade(cell,"1C5A4E")
    rr=cell.paragraphs[0].add_run(c); rr.bold=True; rr.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); rr.font.size=Pt(9)
exemplos = [("1","[Instalar monitoramento (selo de gesso) nas fissuras 1 a 3]","[imediato]","[eng. responsável]","[Alta/Média/Baixa]"),
            ("2","[Investigação geotécnica/de fundações]","[__ dias]","[especialista]","[ ]"),
            ("3","[Tratar origem de umidade antes do acabamento]","[__ dias]","[ ]","[ ]")]
for ex in exemplos:
    cells=t3.add_row().cells
    for i,v in enumerate(ex):
        rr=cells[i].paragraphs[0].add_run(v); rr.font.size=Pt(9); rr.font.color.rgb=REALCE

# 8. Encaminhamentos / conclusão
h("8. ENCAMINHAMENTOS E CONCLUSÃO")
doc.add_paragraph("Síntese objetiva do diagnóstico e do próximo passo. Evite promessas de resultado; "
                  "descreva capacidades e recomendações técnicas.")
campo("Conclusão:","ex.: quadro compatível com recalque diferencial ativo; recomenda-se investigação e monitoramento; "
      "havendo evolução, avaliar escoramento/interdição")

# 9. Assinatura
h("9. RESPONSABILIDADE TÉCNICA")
doc.add_paragraph()
ass = doc.add_paragraph(); ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
ass.add_run("_______________________________________________\n").font.color.rgb = CINZA
ra = ass.add_run("Gustavo Domingos — Engenheiro Civil\nCREA-PR 140.964-D · CREA-SP 5071652757\n")
ra.bold = True; ra.font.color.rgb = VERDE
ass.add_run("Perito judicial atuante em processos do TJ-PR e do TJ-SP  [VALIDAR redação — D1]").font.color.rgb = CINZA
foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = foot.add_run("Escola de Obra · GD Engenharia e Perícia · engenhariagd.com.br"); rf.font.size=Pt(8); rf.font.color.rgb=CINZA

out = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "02-curso-piloto", "kit", "relatorio", "relatorio-blindado-fissuras.docx"))
doc.save(out)
print("RELATORIO DOCX GERADO:", out, f"({os.path.getsize(out)} bytes)")
