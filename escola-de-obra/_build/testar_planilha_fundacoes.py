#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Testa a planilha de fundações: injeta 4 casos, avalia as fórmulas de triagem (lib 'formulas')
e confere Família, Tipos, Alerta e o Dashboard."""
import os, re, sys, shutil, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
import formulas

BASE=os.path.abspath(os.path.join(os.path.dirname(__file__),".."))
SRC=os.path.join(BASE,"04-curso-fundacoes","kit","planilha","planilha-fundacoes.xlsx")
TMP="/tmp/planilha-fundacoes-teste.xlsx"
shutil.copy(SRC,TMP)

casos=[
 # B,C,D,E,F,G, esperado(familia, regex_tipos, regex_alerta)
 ("P1 - firme raso",1.5,10,"leve","Nao","Nao",("RASA",r"Sapata",r"^$")),
 ("P2 - argila mole/edificio",9,10,"alta","Nao","Nao",("PROFUNDA",r"tubulao",r"Carga alta")),
 ("P3 - NA alto+vizinho",9,1,"media","Sim","Nao",("PROFUNDA",r"helice continua/escavada",r"NA alto.*|Vizinhanca")),
 ("P4 - intermediaria",5,10,"media","Nao","Nao",("ZONA CINZENTA",r"projetista",r"Decisao do projetista")),
]
wb=load_workbook(TMP); ws=wb["Seletor"]
for i,(b,c,d,e,f,g,_) in enumerate(casos):
    r=3+i
    ws[f"B{r}"]=b; ws[f"C{r}"]=c; ws[f"D{r}"]=d; ws[f"E{r}"]=e; ws[f"F{r}"]=f; ws[f"G{r}"]=g
wb.save(TMP)

xl=formulas.ExcelModel().loads(TMP).finish(); sol=xl.calculate()
def get(sheet,cell):
    pat=re.compile(r"\]"+sheet.upper()+r"'?\!"+cell.upper()+r"$")
    for k,v in sol.items():
        if pat.search(k):
            try: val=v.value[0,0]
            except Exception:
                try: val=v.value
                except Exception: val=v
            if hasattr(val,"ravel"): val=val.ravel()[0]
            return val
    return None

ok=True
print("="*64); print(" TESTE DA PLANILHA DE FUNDAÇÕES — 4 casos avaliados"); print("="*64)
for i,(b,c,d,e,f,g,esp) in enumerate(casos):
    r=3+i
    fam=str(get("Seletor",f"H{r}")).strip().strip('"')
    tipos=str(get("Seletor",f"I{r}")).strip().strip('"')
    alerta=str(get("Seletor",f"J{r}")).strip().strip('"')
    efam,etip,eal=esp
    if alerta in ("None","0"): alerta=""
    fam_ok=fam==efam
    tip_ok=bool(re.search(etip,tipos,re.I))
    al_ok=bool(re.search(eal,alerta,re.I)) if eal!=r"^$" else (alerta=="")
    passed=fam_ok and tip_ok and al_ok; ok=ok and passed
    print(f"\n{'✔' if passed else '✗'} Caso {i+1}: {b}  (camada={c}m NA={d}m carga={e} viz={f})")
    print(f"   Família:  {fam}  (esp {efam}) {'✓' if fam_ok else '✗'}")
    print(f"   Tipos:    {tipos}  {'✓' if tip_ok else '✗'}")
    print(f"   Alerta:   {alerta or '(vazio)'}  {'✓' if al_ok else '✗'}")

def num(x):
    try: return int(round(float(x)))
    except Exception: return x
cas=num(get("Dashboard","A4")); rasa=num(get("Dashboard","B4")); prof=num(get("Dashboard","C4")); zc=num(get("Dashboard","D4"))
print("\n"+"-"*64); print(" DASHBOARD")
dash_ok=(cas==4 and rasa==1 and prof==2 and zc==1)
print(f"  Casos={cas} (4) · RASA={rasa} (1) · PROFUNDA={prof} (2) · ZONA CINZENTA={zc} (1)  {'✔' if dash_ok else '✗'}")
ok=ok and dash_ok
print("\n"+"="*64); print(" RESULTADO FINAL:", "TODOS OS TESTES PASSARAM ✔" if ok else "HÁ FALHAS ✗"); print("="*64)
sys.exit(0 if ok else 1)
