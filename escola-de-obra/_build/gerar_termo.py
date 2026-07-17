#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Termo de Autorização de Uso de Imagem e Depoimento (.docx) — Etapa 5."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

VERDE = RGBColor(0x0E,0x3A,0x34); DOUR = RGBColor(0xC9,0xA2,0x4B); CINZA = RGBColor(0x6B,0x6B,0x63)
doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("TERMO DE AUTORIZAÇÃO DE USO DE IMAGEM E DEPOIMENTO")
r.bold = True; r.font.size = Pt(14); r.font.color.rgb = VERDE; r.font.name = "Georgia"
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = s.add_run("Escola de Obra · GD Engenharia e Perícia"); rs.font.color.rgb = DOUR; rs.bold = True; rs.font.size = Pt(9)
doc.add_paragraph()

def campo(label):
    p = doc.add_paragraph()
    rr = p.add_run(label + " "); rr.bold = True; rr.font.color.rgb = VERDE
    p.add_run("_______________________________________________")

campo("Nome completo:")
campo("CPF:")
campo("Cidade/UF:")
campo("Contato (e-mail/telefone):")
doc.add_paragraph()

corpo = doc.add_paragraph()
corpo.add_run(
    "Pelo presente termo, autorizo, de forma gratuita e por prazo indeterminado, a Escola de Obra / "
    "GD Engenharia e Perícia a utilizar meu depoimento, nome, cidade e imagem (foto e/ou vídeo) fornecidos "
    "por mim, exclusivamente para fins de divulgação do curso e dos materiais da Escola de Obra, em páginas "
    "de venda, redes sociais, e-mails e materiais institucionais.")
doc.add_paragraph()
for item in [
    "Declaro que meu depoimento é verdadeiro e reflete minha experiência real com o curso.",
    "Estou ciente de que posso revogar esta autorização a qualquer momento, por escrito, para usos futuros.",
    "Autorizo o tratamento dos meus dados pessoais para esta finalidade, nos termos da LGPD (Lei 13.709/2018).",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph()
p = doc.add_paragraph(); p.add_run("Autorizo o uso de: ").bold = True
p.add_run("(  ) depoimento em texto   (  ) foto   (  ) vídeo   (  ) nome e cidade")
doc.add_paragraph(); doc.add_paragraph()

for linha in ["Local e data: __________________________, ______ / ______ / __________",
              "", "_______________________________________________",
              "Assinatura"]:
    pp = doc.add_paragraph(linha); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER

f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = f.add_run("Responsável pela coleta: Gustavo Domingos — Eng. Civil, CREA-PR 140.964-D · engenhariagd.com.br")
rf.font.size = Pt(8); rf.font.color.rgb = CINZA

out = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "03-lancamento", "depoimentos", "termo-autorizacao.docx"))
doc.save(out)
print("TERMO GERADO:", out, f"({os.path.getsize(out)} bytes)")
