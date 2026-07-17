#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Testa a planilha de inspeção: injeta 5 itens e avalia situações + veredito global (lib 'formulas')."""
import os, re, sys, shutil, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
import formulas

BASE=os.path.abspath(os.path.join(os.path.dirname(__file__),".."))
SRC=os.path.join(BASE,"06-curso-inspecao","kit","planilha","planilha-inspecao.xlsx")
TMP="/tmp/planilha-inspecao-teste.xlsx"
shutil.copy(SRC,TMP)

itens=[
 # servico, elemento, item, resultado, reconferido, esperado
 ("Forma","V3","Secao 19x40 conferida","Conforme","", "OK"),
 ("Cobrimento","Laje L2","Faltam espacadores bordo leste","Pendencia","Nao","PENDENTE"),
 ("Escoramento","Laje L2","Sem contraventamento direcao Y","NC critica","Nao","NC CRITICA ABERTA"),
 ("Armacao","V4","Amarracao frouxa - refeita","Pendencia","Sim","RESOLVIDA"),
 ("Embutidos","Laje L2","Eletrica conferida c/ equipe","Conforme","","OK"),
]
wb=load_workbook(TMP); ws=wb["Inspecao"]
for i,(sv,el,it,res,rec,_) in enumerate(itens):
    r=3+i
    ws[f"B{r}"]=sv; ws[f"C{r}"]=el; ws[f"D{r}"]=it; ws[f"E{r}"]=res
    if rec: ws[f"G{r}"]=rec
wb.save(TMP)

xl=formulas.ExcelModel().loads(TMP).finish(); sol=xl.calculate()
def get(sheet,cell):
    pat=re.compile(r"\]"+sheet.upper()+r"'?\!"+cell.upper()+r"$")
    for k,v in sol.items():
        if pat.search(k):
            try: val=v.value[0,0]
            except Exception:
                try: val=v.value
                except Exception: val=v
            if hasattr(val,"ravel"): val=val.ravel()[0]
            return val
    return None

ok=True
print("="*62); print(" TESTE DA PLANILHA DE INSPEÇÃO"); print("="*62)
for i,(sv,el,it,res,rec,esp) in enumerate(itens):
    r=3+i
    sit=str(get("Inspecao",f"H{r}")).strip().strip('"')
    p=sit==esp; ok=ok and p
    print(f"{'✔' if p else '✗'} {sv:12s} {res:10s} rec={rec or '—':3s} → {sit}  (esp {esp})")

def num(x):
    try: return int(round(float(x)))
    except Exception: return x
tot=num(get("Dashboard","A4")); okc=num(get("Dashboard","B4")); pend=num(get("Dashboard","C4"))
nc=num(get("Dashboard","D4")); res=num(get("Dashboard","E4"))
ver=str(get("Dashboard","B6")).strip().strip('"')
print("\n"+"-"*62)
dash_ok=(tot==5 and okc==2 and pend==1 and nc==1 and res==1 and ver.startswith("TRAVADO"))
print(f" DASHBOARD: itens={tot}(5) OK={okc}(2) pend={pend}(1) NC={nc}(1) resolvidas={res}(1)")
print(f" VEREDITO: {ver}  {'✔' if ver.startswith('TRAVADO') else '✗'}")
ok=ok and dash_ok
print("\n"+"="*62); print(" RESULTADO FINAL:", "TODOS OS TESTES PASSARAM ✔" if ok else "HÁ FALHAS ✗"); print("="*62)
sys.exit(0 if ok else 1)
