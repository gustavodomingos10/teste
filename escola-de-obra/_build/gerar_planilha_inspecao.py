#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilha de Inspeção Multi-serviço (Kit do Curso 4).
Abas: Inspecao (uma linha por item verificado, situação automática), Dashboard (veredito global).
Saída: 06-curso-inspecao/kit/planilha/planilha-inspecao.xlsx"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

VERDE="0E3A34"; VMED="1C5A4E"; CREME="F5EFE0"; CREME2="FBF8F0"; DOUR="C9A24B"; CARVAO="20211E"; CINZA="6B6B63"
f_tit=Font(name="Georgia",size=15,bold=True,color=CREME)
f_hdr=Font(name="Calibri",size=10,bold=True,color=CREME)
f_lbl=Font(name="Calibri",size=10,bold=True,color=VERDE)
f_norm=Font(name="Calibri",size=10,color=CARVAO)
f_small=Font(name="Calibri",size=8.5,italic=True,color=CINZA)
fill_v=PatternFill("solid",fgColor=VERDE); fill_vm=PatternFill("solid",fgColor=VMED)
fill_c2=PatternFill("solid",fgColor=CREME2)
fill_ok=PatternFill("solid",fgColor="D9EAD3"); fill_warn=PatternFill("solid",fgColor="FFF2CC")
fill_bad=PatternFill("solid",fgColor="F4CCCC"); fill_res=PatternFill("solid",fgColor="CFE2F3")
thin=Side(style="thin",color="D9D2BE"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal="center",vertical="center",wrap_text=True)
left=Alignment(horizontal="left",vertical="center",wrap_text=True)

wb=Workbook()
ws=wb.active; ws.title="Inspecao"; ws.sheet_view.showGridLines=False
ws["A1"]="INSPEÇÃO DE SERVIÇOS CRÍTICOS — uma linha por item verificado (NBR 14931:2023 · 6118 · 15696 · 8545)"
ws["A1"].font=f_tit; ws["A1"].fill=fill_v; ws.merge_cells("A1:I1"); ws.row_dimensions[1].height=24
ws["A1"].alignment=Alignment(horizontal="left",vertical="center")
headers=[("A","Data",11),("B","Serviço",13),("C","Elemento / Local",18),("D","Item verificado",26),
 ("E","Resultado",13),("F","Prazo correção",12),("G","Reconferido?",11),("H","Situação",16),("I","Observações / foto",24)]
for col,txt,w in headers:
    c=ws[f"{col}2"]; c.value=txt; c.font=f_hdr; c.fill=fill_v; c.alignment=center; c.border=border
    ws.column_dimensions[col].width=w
ws.row_dimensions[2].height=28; ws.freeze_panes="A3"
FIRST,LAST=3,202
for row in range(FIRST,LAST+1):
    ws[f"H{row}"]=(f'=IF(E{row}="","",'
      f'IF(E{row}="Conforme","OK",'
      f'IF(G{row}="Sim","RESOLVIDA",'
      f'IF(E{row}="NC critica","NC CRITICA ABERTA","PENDENTE"))))')
    for col,_,_ in headers:
        cell=ws[f"{col}{row}"]; cell.font=f_norm; cell.border=border
        cell.alignment=center if col in "ABEFGH" else left
        if row%2==0: cell.fill=fill_c2
    ws[f"A{row}"].number_format="dd/mm/yyyy"
dvs={"B":'"Forma,Armacao,Cobrimento,Escoramento,Embutidos,Alvenaria,Desforma"',
     "E":'"Conforme,Pendencia,NC critica"',"G":'"Sim,Nao"'}
for col,formula in dvs.items():
    dv=DataValidation(type="list",formula1=formula,allow_blank=True); dv.showInputMessage=True
    dv.prompt="Selecione"; ws.add_data_validation(dv); dv.add(f"{col}{FIRST}:{col}{LAST}")
rng=f"H{FIRST}:H{LAST}"
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"OK"'],fill=fill_ok,font=Font(color="1b4d1f",bold=True)))
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"PENDENTE"'],fill=fill_warn,font=Font(color="9C6500",bold=True)))
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"NC CRITICA ABERTA"'],fill=fill_bad,font=Font(color="7f1d1d",bold=True)))
ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=['"RESOLVIDA"'],fill=fill_res,font=Font(color="0b3d66",bold=True)))

wd=wb.create_sheet("Dashboard"); wd.sheet_view.showGridLines=False
wd["A1"]="DASHBOARD — Veredito da liberação"; wd["A1"].font=f_tit; wd["A1"].fill=fill_v
wd.merge_cells("A1:E1"); wd.row_dimensions[1].height=26; wd["A1"].alignment=Alignment(horizontal="left",vertical="center")
def metric(lc,vc,label,formula,fill):
    wd[lc]=label; wd[lc].font=f_lbl; wd[lc].border=border; wd[lc].alignment=left
    wd[vc]=formula; wd[vc].font=Font(name="Georgia",size=18,bold=True,color=VERDE); wd[vc].alignment=center; wd[vc].fill=fill; wd[vc].border=border
metric("A3","A4","Itens verificados",'=COUNTA(Inspecao!D3:D202)',fill_c2)
metric("B3","B4","OK",'=COUNTIF(Inspecao!H3:H202,"OK")',fill_ok)
metric("C3","C4","Pendentes",'=COUNTIF(Inspecao!H3:H202,"PENDENTE")',fill_warn)
metric("D3","D4","NC críticas abertas",'=COUNTIF(Inspecao!H3:H202,"NC CRITICA ABERTA")',fill_bad)
metric("E3","E4","Resolvidas",'=COUNTIF(Inspecao!H3:H202,"RESOLVIDA")',fill_res)
wd["A6"]="VEREDITO"; wd["A6"].font=f_hdr; wd["A6"].fill=fill_vm; wd["A6"].border=border; wd["A6"].alignment=center
wd["B6"]=('=IF(D4>0,"TRAVADO - corrigir NC criticas e reinspecionar",'
 'IF(C4>0,"LIBERADO COM PENDENCIAS - prazo + reconferencia obrigatorios",'
 'IF(B4+E4>0,"LIBERADO - preencher o Termo e reconferir no dia","")))')
wd.merge_cells("B6:E6"); wd["B6"].font=Font(name="Calibri",size=12,bold=True,color=VERDE)
wd["B6"].alignment=left; wd["B6"].border=border
wd["A8"]=('Como usar: uma linha por item verificado (serviço, elemento, item, resultado). Pendência/NC '
 'corrigida e reconferida vira RESOLVIDA. O veredito global exige zero NC crítica aberta. '
 'Este controle apoia o Termo de Liberação — não substitui o projeto nem o responsável técnico.')
wd["A8"].font=f_small; wd.merge_cells("A8:E10"); wd["A8"].alignment=left
wd["A12"]="Revisado e assinado por Gustavo Domingos — Engenheiro Civil, CREA-PR 140.964-D"
wd["A12"].font=Font(name="Calibri",size=9,bold=True,color=DOUR)
for w,c in [(18,"A"),(12,"B"),(12,"C"),(17,"D"),(12,"E")]: wd.column_dimensions[c].width=w

out=os.path.abspath(os.path.join(os.path.dirname(__file__),"..","06-curso-inspecao","kit","planilha","planilha-inspecao.xlsx"))
wb.save(out); print("PLANILHA GERADA:",out)
